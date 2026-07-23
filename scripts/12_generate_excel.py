"""
STEP 12 - GENERATE EXCEL

Input   : data/corrected_json/{pdf_name}/page_N_corrected.json (all pages)
          constituency_config/{pdf_name}.json (candidate names for headers)
Output  : data/excel_reports/{pdf_name}/{pdf_name}_final.xlsx

Reads every field from the corrected JSON and writes it into a single
Excel sheet: Page, Booth ID, one column per candidate (using real names),
Total Valid Votes, Rejected Votes, NOTA, Total Votes, Tendered Votes,
Validation Status.
"""

import json
import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from config import CORRECTED_JSON_DIR, EXCEL_REPORTS_DIR, PDF_NAMES, load_constituency_config

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REVIEW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def process_pdf(pdf_name: str):
    src_dir = CORRECTED_JSON_DIR / pdf_name
    out_dir = EXCEL_REPORTS_DIR / pdf_name

    if not src_dir.exists():
        print(f"[{pdf_name}] SKIP - run 11_final_validation.py first")
        return

    out_dir.mkdir(parents=True, exist_ok=True)

    candidates = load_constituency_config(pdf_name)["candidates"]

    wb = Workbook()
    ws = wb.active
    ws.title = pdf_name[:31]

    headers = (
        ["Page", "Booth ID"]
        + candidates
        + ["Total Valid Votes", "Rejected Votes", "NOTA", "Total Votes",
           "Tendered Votes", "Validation Status"]
    )
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    row_count = 0
    for path in sorted(src_dir.glob("page_*_corrected.json"),
                        key=lambda p: int(p.stem.split("_")[1])):
        with open(path, "r", encoding="utf-8") as f:
            page_data = json.load(f)

        for booth in page_data["booths"]:
            status = booth.get("validation_status", "unknown")
            row = (
                [page_data["page"], booth["booth_id"]]
                + booth["candidate_votes"]
                + [
                    booth["total_valid_votes"],
                    booth["rejected_votes"],
                    booth["nota"],
                    booth["total_votes"],
                    booth["tendered_votes"],
                    status,
                ]
            )
            ws.append(row)
            row_count += 1

            if status == "manual_review":
                for cell in ws[ws.max_row]:
                    cell.fill = REVIEW_FILL

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)

    out_path = out_dir / f"{pdf_name}_final.xlsx"
    wb.save(out_path)
    print(f"[{pdf_name}] Saved {row_count} booth rows -> {out_path}")


if __name__ == "__main__":
    for pdf_name in PDF_NAMES:
        process_pdf(pdf_name)