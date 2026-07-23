"""
STEP 13 - MASTER EXCEL (single sheet per PDF: data + summary combined)

Input   : data/corrected_json/{pdf_name}/page_N_corrected.json (all PDFs,
          including 1-Sirpur, 2-Chennur, 3-Bellampalli from batch 1)
          constituency_config/{pdf_name}.json (candidate names)
Output  : data/master_report/All_Constituencies_Master.xlsx

For each PDF, creates ONE sheet in the master workbook containing:
  1. Every booth row (Page, Booth ID, candidate votes, totals, status)
  2. Followed immediately below by the summary/analysis section:
     overall stats, winning candidate, runner-up, victory margin,
     highest/lowest polling booths, candidate-wise totals, page-wise
     breakdown.

FIX: previous version always created a brand new empty workbook and only
looped over PDF_NAMES (the current config list) - this silently
overwrote the master file and lost batch 1's sheets (Sirpur, Chennur,
Bellampalli) since they aren't in PDF_NAMES. This version now:
  1. Loads the EXISTING master file if present, instead of starting fresh
  2. Explicitly includes the batch 1 PDFs alongside PDF_NAMES
  3. Skips any PDF whose sheet is already present, so re-running this
     safely does nothing destructive to sheets that already exist
"""

import json
import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from config import CORRECTED_JSON_DIR, MASTER_REPORT_DIR, PDF_NAMES, load_constituency_config

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUMMARY_LABEL_FONT = Font(bold=True)
SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

# Batch 1 PDFs - not in PDF_NAMES (which only lists the current batch),
# so they must be added explicitly to avoid losing them from the master file.
BATCH_1_PDFS = ["1-Sirpur", "2-Chennur", "3-Bellampalli"]


def load_all_pages(pdf_name: str):
    src_dir = CORRECTED_JSON_DIR / pdf_name
    pages = []
    if not src_dir.exists():
        return pages
    for path in sorted(src_dir.glob("page_*_corrected.json"),
                        key=lambda p: int(p.stem.split("_")[1])):
        with open(path, "r", encoding="utf-8") as f:
            pages.append(json.load(f))
    return pages


def compute_page_stats(page_data):
    booths = page_data["booths"]
    total = len(booths)
    passed = sum(1 for b in booths if b.get("validation_status") != "manual_review")
    failed = total - passed
    accuracy = (passed / total * 100) if total else 0

    total_valid = sum(b["total_valid_votes"] for b in booths)
    total_votes = sum(b["total_votes"] for b in booths)

    highest = max(booths, key=lambda b: b["total_votes"]) if booths else None
    lowest = min(booths, key=lambda b: b["total_votes"]) if booths else None

    return {
        "page": page_data["page"], "total": total, "passed": passed, "failed": failed,
        "accuracy": accuracy, "total_valid": total_valid, "total_votes": total_votes,
        "highest": highest, "lowest": lowest,
    }


def write_combined_sheet(wb, pdf_name, pages, candidates):
    ws = wb.create_sheet(title=pdf_name[:31])

    # ---------- SECTION 1: BOOTH-LEVEL DATA ----------
    headers = (["Page", "Booth ID"] + candidates +
               ["Total Valid Votes", "Rejected Votes", "NOTA", "Total Votes",
                "Tendered Votes", "Validation Status"])
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for page_data in pages:
        for booth in page_data["booths"]:
            ws.append(
                [page_data["page"], booth["booth_id"]]
                + booth["candidate_votes"]
                + [booth["total_valid_votes"], booth["rejected_votes"], booth["nota"],
                   booth["total_votes"], booth["tendered_votes"],
                   booth.get("validation_status", "unknown")]
            )

    # ---------- SECTION 2: SUMMARY / ANALYSIS (appended below) ----------
    all_booths = [b for page_data in pages for b in page_data["booths"]]
    total_booths = len(all_booths)
    passed_booths = sum(1 for b in all_booths if b.get("validation_status") != "manual_review")
    failed_booths = total_booths - passed_booths
    overall_accuracy = (passed_booths / total_booths * 100) if total_booths else 0

    total_valid = sum(b["total_valid_votes"] for b in all_booths)
    total_rejected = sum(b["rejected_votes"] for b in all_booths)
    total_nota = sum(b["nota"] for b in all_booths)
    total_votes = sum(b["total_votes"] for b in all_booths)
    total_tendered = sum(b["tendered_votes"] for b in all_booths)

    candidate_totals = {name: 0 for name in candidates}
    for b in all_booths:
        for name, votes in zip(candidates, b["candidate_votes"]):
            candidate_totals[name] += votes

    ranked = sorted(candidate_totals.items(), key=lambda x: x[1], reverse=True)
    winner_name, winner_votes = ranked[0]
    runner_name, runner_votes = ranked[1] if len(ranked) > 1 else ("N/A", 0)
    margin = winner_votes - runner_votes

    highest_booth = max(all_booths, key=lambda b: b["total_votes"]) if all_booths else None
    lowest_booth = min(all_booths, key=lambda b: b["total_votes"]) if all_booths else None

    def add_row(label, value):
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = SUMMARY_LABEL_FONT

    def add_section(title):
        ws.append([])
        ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).fill = SECTION_FILL
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)

    add_section(f"{pdf_name} - OVERALL SUMMARY")
    add_row("TOTAL BOOTHS", total_booths)
    add_row("PASSED BOOTHS", passed_booths)
    add_row("FAILED BOOTHS", failed_booths)
    add_row("VALIDATION ACCURACY", f"{overall_accuracy:.2f}%")
    add_row("TOTAL VALID VOTES", total_valid)
    add_row("TOTAL REJECTED VOTES", total_rejected)
    add_row("TOTAL NOTA", total_nota)
    add_row("TOTAL VOTES", total_votes)
    add_row("TOTAL TENDERED VOTES", total_tendered)
    add_row("WINNING CANDIDATE", winner_name)
    add_row("WINNING VOTES", winner_votes)
    add_row("RUNNER UP", runner_name)
    add_row("RUNNER UP VOTES", runner_votes)
    add_row("VICTORY MARGIN", margin)
    if highest_booth:
        add_row("HIGHEST POLLING BOOTH", highest_booth["booth_id"])
        add_row("HIGHEST TOTAL VOTES", highest_booth["total_votes"])
    if lowest_booth:
        add_row("LOWEST POLLING BOOTH", lowest_booth["booth_id"])
        add_row("LOWEST TOTAL VOTES", lowest_booth["total_votes"])

    add_section("CANDIDATE-WISE TOTAL VOTES")
    ws.append(["CANDIDATE", "TOTAL VOTES"])
    for cell in ws[ws.max_row]:
        cell.font = SUMMARY_LABEL_FONT
    for name, votes in ranked:
        ws.append([name, votes])

    add_section("PAGE-WISE BREAKDOWN")
    ws.append(["PAGE", "TOTAL BOOTHS", "PASSED", "FAILED", "ACCURACY",
               "TOTAL VALID VOTES", "TOTAL VOTES", "HIGHEST BOOTH",
               "HIGHEST VOTES", "LOWEST BOOTH", "LOWEST VOTES"])
    for cell in ws[ws.max_row]:
        cell.font = SUMMARY_LABEL_FONT
    for page_data in pages:
        stats = compute_page_stats(page_data)
        ws.append([
            stats["page"], stats["total"], stats["passed"], stats["failed"],
            f"{stats['accuracy']:.2f}%", stats["total_valid"], stats["total_votes"],
            stats["highest"]["booth_id"] if stats["highest"] else "N/A",
            stats["highest"]["total_votes"] if stats["highest"] else "N/A",
            stats["lowest"]["booth_id"] if stats["lowest"] else "N/A",
            stats["lowest"]["total_votes"] if stats["lowest"] else "N/A",
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 35)


def main():
    MASTER_REPORT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = MASTER_REPORT_DIR / "All_Constituencies_Master.xlsx"

    if out_path.exists():
        wb = load_workbook(out_path)
        print(f"Loaded EXISTING master file. Current sheets: {wb.sheetnames}")
    else:
        wb = Workbook()
        wb.remove(wb.active)
        print("No existing master file found - creating a new one.")

    # Batch 1 PDFs first (in case they're missing and need to be added back),
    # then everything currently in PDF_NAMES.
    all_pdfs = BATCH_1_PDFS + list(PDF_NAMES)

    for pdf_name in all_pdfs:
        sheet_name = pdf_name[:31]
        if sheet_name in wb.sheetnames:
            print(f"[{pdf_name}] SKIP - sheet already exists in master file, not overwritten")
            continue

        pages = load_all_pages(pdf_name)
        if not pages:
            print(f"[{pdf_name}] SKIP - no corrected_json found, run 10/11 first")
            continue

        candidates = load_constituency_config(pdf_name)["candidates"]
        write_combined_sheet(wb, pdf_name, pages, candidates)
        print(f"[{pdf_name}] Added combined sheet (data + summary)")

    wb.save(out_path)
    print(f"\nSaved master report: {out_path}")
    print(f"Final sheets in workbook: {wb.sheetnames}")


if __name__ == "__main__":
    main()